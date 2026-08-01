from __future__ import annotations

import csv
import io
import re
from datetime import datetime, time as dt_time, timedelta, timezone

from src.ingestion.models import AssetRecord

VALID_ASSET_TYPES = {"episode", "slate", "bumper"}
REQUIRED_HEADERS = {"asset_id", "asset_type", "duration_hms"}

_HMS_RE = re.compile(r"^(\d+):([0-5]?\d):([0-5]?\d)(?:\.(\d+))?$")


def _parse_hms_to_ms(value: str) -> int | None:
    value = value.strip()
    if not value:
        return None
    m = _HMS_RE.match(value)
    if not m:
        return None
    hours, minutes, seconds, frac = m.groups()
    total_ms = (int(hours) * 3600 + int(minutes) * 60 + int(seconds)) * 1000
    if frac:
        total_ms += round(float(f"0.{frac}") * 1000)
    return total_ms


def _duration_value_to_ms(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt_time):
        return ((value.hour * 60 + value.minute) * 60 + value.second) * 1000 + value.microsecond // 1000
    if isinstance(value, timedelta):
        return int(value.total_seconds() * 1000)
    if isinstance(value, str):
        return _parse_hms_to_ms(value)
    return None


def _parse_iso_dt(value: str) -> datetime | None:
    value = value.strip()
    if not value:
        return None
    cleaned = value.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(cleaned)
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _valid_value_to_dt(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, str):
        return _parse_iso_dt(value)
    return None


def _str_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _int_or_none(value: object, field: str, row_num: int, errors: list[str]) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        errors.append(f"row {row_num}: {field} must be a number")
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        errors.append(f"row {row_num}: {field} '{value}' is not a valid number")
        return None


def _build_segments_from_cue_points(cue_points_ms: list[int], duration_ms: int) -> list[dict]:
    """Reconstruct the segment list the scheduler expects (see scheduler/service.py
    _build_cue_points_by_asset) from a flat list of absolute ms offsets."""
    segments: list[dict] = []
    prev = 0
    for order, cue in enumerate(cue_points_ms, start=1):
        segments.append({"order": order, "duration_ms": cue - prev, "insert_ad_break": True})
        prev = cue
    remainder = duration_ms - prev
    if remainder > 0:
        segments.append({"order": len(cue_points_ms) + 1, "duration_ms": remainder, "insert_ad_break": False})
    return segments


def _row_to_asset(row: dict[str, object], row_num: int) -> tuple[AssetRecord | None, list[str]]:
    errors: list[str] = []

    asset_id = _str_or_none(row.get("asset_id"))
    if not asset_id:
        return None, [f"row {row_num}: asset_id is required"]

    asset_type_raw = _str_or_none(row.get("asset_type"))
    asset_type = asset_type_raw.lower() if asset_type_raw else None
    if asset_type not in VALID_ASSET_TYPES:
        return None, [f"row {row_num} ({asset_id}): asset_type must be one of episode/slate/bumper"]

    duration_ms = _duration_value_to_ms(row.get("duration_hms"))
    if not duration_ms or duration_ms <= 0:
        return None, [f"row {row_num} ({asset_id}): duration_hms is required and must be > 0 (HH:MM:SS)"]

    season_number = _int_or_none(row.get("season_number"), "season_number", row_num, errors)
    episode_number = _int_or_none(row.get("episode_number"), "episode_number", row_num, errors)

    valid_from_raw = row.get("valid_from")
    valid_to_raw = row.get("valid_to")
    valid_from = _valid_value_to_dt(valid_from_raw)
    valid_to = _valid_value_to_dt(valid_to_raw)
    if valid_from is None and valid_from_raw not in (None, ""):
        errors.append(f"row {row_num} ({asset_id}): valid_from is not a valid ISO 8601 datetime")
    if valid_to is None and valid_to_raw not in (None, ""):
        errors.append(f"row {row_num} ({asset_id}): valid_to is not a valid ISO 8601 datetime")
    if valid_from and valid_to and valid_to <= valid_from:
        errors.append(f"row {row_num} ({asset_id}): valid_to must be after valid_from")

    cue_points_raw = row.get("ad_cue_points_hms")
    cue_points_ms: list[int] = []
    if cue_points_raw not in (None, ""):
        for part in str(cue_points_raw).split("|"):
            part = part.strip()
            if not part:
                continue
            ms = _parse_hms_to_ms(part)
            if ms is None:
                errors.append(f"row {row_num} ({asset_id}): ad_cue_points_hms segment '{part}' is not valid HH:MM:SS")
                continue
            cue_points_ms.append(ms)
        cue_points_ms.sort()
        if cue_points_ms and cue_points_ms[-1] >= duration_ms:
            errors.append(f"row {row_num} ({asset_id}): ad_cue_points_hms contains an offset at/beyond duration_hms")
            cue_points_ms = [c for c in cue_points_ms if c < duration_ms]

    if errors:
        return None, errors

    segments = _build_segments_from_cue_points(cue_points_ms, duration_ms) if cue_points_ms else []

    asset = AssetRecord(
        asset_id=asset_id,
        asset_type=asset_type,
        title=_str_or_none(row.get("title")),
        description=_str_or_none(row.get("description")),
        rating=_str_or_none(row.get("rating")),
        genre=_str_or_none(row.get("genre")),
        tms_id=_str_or_none(row.get("tms_id")),
        series_id=None,
        season_id=None,
        season_number=season_number,
        episode_number=episode_number,
        thumbnail_url=_str_or_none(row.get("thumbnail_url")),
        subtitle_url=_str_or_none(row.get("subtitle_url")),
        valid_from=valid_from,
        valid_to=valid_to,
        duration_ms=duration_ms,
        raw_payload={"source": "spreadsheet", "segments": segments},
    )
    return asset, []


def _parse_rows(rows: list[dict[str, object]], header_names: set[str]) -> tuple[list[AssetRecord], list[str]]:
    missing = REQUIRED_HEADERS - header_names
    if missing:
        return [], [f"missing required column(s): {', '.join(sorted(missing))}"]

    assets: list[AssetRecord] = []
    errors: list[str] = []
    for row_num, row in enumerate(rows, start=2):  # row 1 is the header
        asset, row_errors = _row_to_asset(row, row_num)
        if asset is not None:
            assets.append(asset)
        errors.extend(row_errors)
    return assets, errors


def parse_csv_bytes(data: bytes) -> tuple[list[AssetRecord], list[str]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return [], ["CSV file has no header row"]
    key_map = {name: name.strip().lower() for name in reader.fieldnames if name}
    rows = [
        {key_map.get(k, (k or "").strip().lower()): v for k, v in raw_row.items() if k}
        for raw_row in reader
    ]
    return _parse_rows(rows, set(key_map.values()))


def parse_xlsx_bytes(data: bytes) -> tuple[list[AssetRecord], list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["Spreadsheet has no header row"]

    key_map = {i: (str(h).strip().lower() if h is not None else "") for i, h in enumerate(header_row)}
    rows = []
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        rows.append({key_map.get(i, ""): v for i, v in enumerate(row) if key_map.get(i)})
    return _parse_rows(rows, set(key_map.values()))


def parse_spreadsheet(data: bytes, filename: str) -> tuple[list[AssetRecord], list[str]]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        return parse_csv_bytes(data)
    if ext in ("xlsx", "xlsm"):
        return parse_xlsx_bytes(data)
    return [], [f"Unsupported file type '.{ext}'. Expected .csv or .xlsx"]
